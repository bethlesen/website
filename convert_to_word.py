#!/usr/bin/env python3

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

# Read the markdown file
with open('hidden-curriculum-job-searching.md', 'r') as f:
    content = f.read()

# Create a new Word document
doc = Document()

# Set up default styles
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Split content into lines
lines = content.split('\n')

for line in lines:
    line = line.strip()

    if not line:
        continue

    # Handle H1 (# Title)
    if line.startswith('# '):
        title = line[2:]
        heading = doc.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Handle H2 (## Section)
    elif line.startswith('## '):
        section = line[3:]
        heading = doc.add_heading(section, level=2)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Handle regular paragraphs
    else:
        # Handle bold text (**text**)
        paragraph = doc.add_paragraph()

        # Simple regex to find bold text
        parts = re.split(r'\*\*(.+?)\*\*', line)

        for i, part in enumerate(parts):
            if i % 2 == 0:  # Regular text
                if part:
                    paragraph.add_run(part)
            else:  # Bold text
                paragraph.add_run(part).bold = True

# Save the document
doc.save('hidden-curriculum-job-searching.docx')
print("Word document created successfully: hidden-curriculum-job-searching.docx")
